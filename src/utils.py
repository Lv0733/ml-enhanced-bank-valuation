"""
utils.py -- Enhanced Recursive LightGBM RIV V4.2
=====================================================
Every function in this file is copied, formula-for-formula, from the
verified, working enhanced_riv_v4.py source code. This file was built by
direct inspection of that source (not from memory or from any prompt's
prose description of V4), specifically to avoid repeating the mistake
documented in the V4.1 debugging report (where a feature set and a
terminal value formula were reconstructed incorrectly from memory).

THE ONLY FUNCTION IN THIS ENTIRE FILE THAT BEHAVES DIFFERENTLY FROM V4 IS
`forecast_roe_spread_path()`, and even there, the only difference is a
single line in the Q9-Q20 phase:
    V4:    omega_final = omega_estimated
    V4.2:  omega_final = omega_estimated if omega_cap is None else min(omega_estimated, omega_cap)
An `omega_cap` parameter (default None) controls this -- passing
omega_cap=None reproduces V4's own unconstrained formula EXACTLY (used
internally by the validation suite to prove self-consistency); passing
omega_cap=0.92 (config.OMEGA_CAP) gives V4.2's constrained behaviour.

A SECOND, EQUALLY IMPORTANT correction versus an earlier (incorrect)
standalone attempt: V4's actual `compute_terminal_value()` formula is
    TV_at_horizon = RI_terminal * (1 + g) / (CoE_q - g),   g = 0.0 (default)
i.e. a FLAT (no-growth) perpetuity that does NOT depend on omega at all.
omega's only effect on terminal value is INDIRECT: it changes the Q9-Q20
spread path, which changes ROE_20 and hence RI_20 (the input to this
formula), but the formula itself never references omega. An earlier
implementation attempt used a different, omega-dependent Ohlson-style
formula here, which was WRONG relative to V4's actual code -- verified by
direct inspection of enhanced_riv_v4.py's compute_terminal_value()
function, reproduced exactly below.
"""

import re
import numpy as np
import pandas as pd
import openpyxl
import lightgbm as lgb
from sklearn.preprocessing import StandardScaler, OneHotEncoder
from sklearn.compose import ColumnTransformer

import config as cfg


# ============================================================================
# DATA LOADING -- identical to V4's load_jpm_financials / load_jpm_weekly_prices
# / load_sp500_daily / load_tbill_quarterly / build_return_panel
# ============================================================================

def _get_row(ws, row_num):
    return list(ws.iter_rows(min_row=row_num, max_row=row_num, values_only=True))[0]


def load_jpm_financials(path: str = None) -> pd.DataFrame:
    path = path or cfg.JPM_FINANCIALS_PATH
    wb = openpyxl.load_workbook(path, data_only=True)
    ws_fs, ws_is, ws_bs = wb["Financial Summary"], wb["Income Statement"], wb["Balance Sheet"]
    dates = _get_row(ws_fs, 12)[1:]
    n = len(dates)
    df = pd.DataFrame({
        "period_end": dates,
        "net_income_avail_common": _get_row(ws_is, 90)[1:n + 1],
        "common_equity": _get_row(ws_bs, 99)[1:n + 1],
        "shares_outstanding_mm": _get_row(ws_bs, 106)[1:n + 1],
        "dps": _get_row(ws_is, 132)[1:n + 1],
        "bvps": _get_row(ws_fs, 41)[1:n + 1],
    })
    df["period_end"] = pd.to_datetime(df["period_end"])
    df = df.dropna(subset=["net_income_avail_common", "common_equity", "shares_outstanding_mm"])
    df = df.sort_values("period_end").reset_index(drop=True)
    df["avg_common_equity"] = (df["common_equity"] + df["common_equity"].shift(1)) / 2
    df["roe_quarterly_annualized_pct"] = (df["net_income_avail_common"] / df["avg_common_equity"]) * 4 * 100
    return df.reset_index(drop=True)


def load_jpm_weekly_prices(path: str = None) -> pd.DataFrame:
    path = path or cfg.JPM_PRICE_PATH
    df = pd.read_csv(path)
    df["Date"] = pd.to_datetime(df["Date"], format="%m/%d/%Y")
    df = df[["Date", "Price"]].rename(columns={"Price": "jpm_price"})
    df = df.sort_values("Date").reset_index(drop=True)
    df["week_end_date"] = df["Date"] + pd.Timedelta(days=5)
    return df


def load_sp500_daily(path: str = None) -> pd.DataFrame:
    path = path or cfg.SP500_PATH
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Daily, Close"]
    rows = list(ws.iter_rows(values_only=True))[1:]
    df = pd.DataFrame(rows, columns=["Date", "sp500"])
    df["Date"] = pd.to_datetime(df["Date"])
    return df.dropna().sort_values("Date").reset_index(drop=True)


def load_tbill_quarterly(path: str = None) -> pd.DataFrame:
    path = path or cfg.TBILL_PATH
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Quarterly"]
    rows = list(ws.iter_rows(values_only=True))[1:]
    df = pd.DataFrame(rows, columns=["Date", "DTB3"])
    df["Date"] = pd.to_datetime(df["Date"])
    return df.dropna().sort_values("Date").reset_index(drop=True)


def build_return_panel(jpm_prices: pd.DataFrame, sp500: pd.DataFrame) -> pd.DataFrame:
    merged = pd.merge_asof(jpm_prices, sp500, left_on="week_end_date", right_on="Date",
                            direction="backward", suffixes=("_jpm", "_sp"))
    merged["jpm_ret"] = merged["jpm_price"].pct_change()
    merged["sp_ret"] = merged["sp500"].pct_change()
    return merged.dropna(subset=["jpm_ret", "sp_ret"]).reset_index(drop=True)[
        ["week_end_date", "jpm_price", "sp500", "jpm_ret", "sp_ret"]]


def quarter_to_end_date(q: str) -> pd.Timestamp:
    year, qtr = int(q[:4]), int(q[-1])
    month = qtr * 3
    return pd.Timestamp(year=year, month=month, day=1) + pd.offsets.MonthEnd(0)


def quarter_add(q: str, n: int) -> str:
    year, qtr = int(q[:4]), int(q[-1])
    total = year * 4 + (qtr - 1) + n
    return f"{total // 4}Q{total % 4 + 1}"


# ============================================================================
# COST OF EQUITY -- identical to V4's get_risk_free_rate / get_erp /
# estimate_beta_expanding / build_historical_coe_series
# ============================================================================

def get_risk_free_rate(tbill_df: pd.DataFrame, as_of: pd.Timestamp) -> float:
    avail = tbill_df[tbill_df["Date"] <= as_of]
    if avail.empty:
        return np.nan
    return avail.iloc[-1]["DTB3"] / 100.0


def get_erp(as_of: pd.Timestamp) -> float:
    return cfg.DAMODARAN_ERP_BY_YEAR.get(as_of.year, np.nan)


def estimate_beta_expanding(return_panel: pd.DataFrame, as_of: pd.Timestamp, min_weeks: int = 104):
    window = return_panel[return_panel["week_end_date"] <= as_of]
    if len(window) < min_weeks:
        return np.nan, len(window)
    cov = np.cov(window["jpm_ret"], window["sp_ret"])
    return cov[0, 1] / cov[1, 1], len(window)


def estimate_cost_of_equity(tbill_df: pd.DataFrame, return_panel: pd.DataFrame, as_of: pd.Timestamp) -> dict:
    """Convenience wrapper bundling the three CAPM components (identical
    formula to V4's inline construction: CoE = Rf + Beta x ERP)."""
    rf = get_risk_free_rate(tbill_df, as_of)
    erp = get_erp(as_of)
    beta, n_weeks = estimate_beta_expanding(return_panel, as_of)
    coe = rf + beta * erp if not (np.isnan(rf) or np.isnan(beta) or np.isnan(erp)) else np.nan
    return {"risk_free_rate": rf, "beta": beta, "beta_n_weeks": n_weeks, "equity_risk_premium": erp, "cost_of_equity": coe}


def build_historical_coe_series(tbill: pd.DataFrame, return_panel: pd.DataFrame, quarters: list) -> pd.DataFrame:
    """Builds CoE at EVERY historical quarter (not just the 9 valuation
    dates), expanding-window, no look-ahead. Identical to V4."""
    rows = []
    for q in quarters:
        as_of = quarter_to_end_date(q)
        rf = get_risk_free_rate(tbill, as_of)
        erp = get_erp(as_of)
        beta, n_weeks = estimate_beta_expanding(return_panel, as_of)
        coe = rf + beta * erp if not (np.isnan(rf) or np.isnan(beta) or np.isnan(erp)) else np.nan
        rows.append({"Quarter": q, "quarter_end": as_of, "risk_free_rate": rf, "beta": beta,
                      "beta_n_weeks": n_weeks, "erp": erp, "cost_of_equity": coe})
    return pd.DataFrame(rows)


# ============================================================================
# LARGE-BANK TRAINING PANEL -- identical to V4's build_large_bank_panel
# (whole-word matching, NOT naive substring matching)
# ============================================================================

def build_large_bank_panel(raw_path: str = None) -> tuple:
    raw_path = raw_path or cfg.ML_PANEL_PATH
    raw = pd.read_csv(raw_path)
    actual_banks = raw["Bank"].unique().tolist()
    actual_words = {b: set(re.findall(r"[a-z]+", b.lower())) for b in actual_banks}

    match_report = []
    matched_actual_names = []
    for requested in cfg.LARGE_BANKS_REQUESTED:
        req_words = set(re.findall(r"[a-z]+", requested.lower()))
        hits = [b for b in actual_banks if req_words.issubset(actual_words[b])]
        found = len(hits) > 0
        match_report.append({"requested_bank": requested, "found_in_dataset": found,
                              "matched_name": hits[0] if found else None})
        if found:
            matched_actual_names.append(hits[0])

    match_df = pd.DataFrame(match_report)
    panel = raw[raw["Bank"].isin(matched_actual_names)].copy()
    panel = panel.sort_values(["Bank", "Quarter"]).reset_index(drop=True)
    return panel, match_df


# ============================================================================
# FEATURE ENGINEERING -- identical to V4's engineer_features(). This is the
# exact function whose feature list was reconstructed incorrectly in the
# prior standalone attempt; it is now copied verbatim from the verified
# source rather than reconstructed.
# ============================================================================

def engineer_features(panel: pd.DataFrame, coe_hist: pd.DataFrame) -> pd.DataFrame:
    """
    Builds ROE (single-quarter-annualised), ROE_SPREAD = ROE - CoE, every
    lag/rolling-window/spread feature, and LLP_to_Assets (replacing raw
    LLP). All rolling/lag features are computed PER BANK via groupby, with
    shift(1) applied BEFORE any rolling window, so no feature ever
    includes the current observation (no look-ahead).
    """
    df = panel.copy().sort_values(["Bank", "Quarter"]).reset_index(drop=True)

    df["avg_common_equity_v4"] = df.groupby("Bank")["Common Equity - Total"].transform(lambda x: (x + x.shift(1)) / 2)
    df["ROE"] = (df["Net Income"] / df["avg_common_equity_v4"]) * 4 * 100  # % units

    df["LLP_to_Assets"] = df["Provision & Impairment for Loan Losses (LLP)"] / df["Total Assets"] * 100

    coe_map = coe_hist.set_index("Quarter")["cost_of_equity"]
    df["CostOfEquity"] = df["Quarter"].map(coe_map) * 100  # % units, matching ROE's % convention

    df["ROE_SPREAD"] = df["ROE"] - df["CostOfEquity"]

    g = df.groupby("Bank")
    df["ROE_Lag1"] = g["ROE"].shift(1)
    df["ROE_Lag2"] = g["ROE"].shift(2)
    df["ROE_Lag4"] = g["ROE"].shift(4)
    df["ROE_MA_4Q"] = g["ROE"].transform(lambda x: x.shift(1).rolling(4, min_periods=4).mean())
    df["ROE_MA_8Q"] = g["ROE"].transform(lambda x: x.shift(1).rolling(8, min_periods=8).mean())
    df["ROE_STD_4Q"] = g["ROE"].transform(lambda x: x.shift(1).rolling(4, min_periods=4).std())

    df["ROE_SPREAD_Lag1"] = g["ROE_SPREAD"].shift(1)
    df["ROE_SPREAD_Lag2"] = g["ROE_SPREAD"].shift(2)
    df["ROE_SPREAD_MA_4Q"] = g["ROE_SPREAD"].transform(lambda x: x.shift(1).rolling(4, min_periods=4).mean())

    # Direct multi-horizon targets: ROE_SPREAD at t+1..t+4, per bank
    for h in range(1, cfg.DIRECT_HORIZON_QUARTERS + 1):
        df[f"ROE_SPREAD_t_plus_{h}"] = g["ROE_SPREAD"].shift(-h)

    return df


# ============================================================================
# DYNAMIC BANK-VARIABLE FORECASTING -- identical to V4's
# train_dynamic_feature_models / estimate_macro_ar1 / project_macro /
# forecast_dynamic_features_path
# ============================================================================

def train_dynamic_feature_models(feat_df: pd.DataFrame, valuation_date_quarter: str) -> dict:
    """Trains one small LightGBM model per dynamic bank-quality variable,
    on the expanding (no-look-ahead) panel of all large banks, pooled."""
    def qkey(q):
        y, qq = int(q[:4]), int(q[-1])
        return y * 4 + qq
    cutoff = qkey(valuation_date_quarter)

    models = {}
    for var in cfg.DYNAMIC_FEATURES:
        df = feat_df.copy()
        g = df.groupby("Bank")
        df[f"{var}_Lag1"] = g[var].shift(1)
        df[f"{var}_target"] = df[var]
        df[f"{var}_Lag2"] = g[var].shift(2)
        train_rows = df[df["Quarter"].apply(qkey) < cutoff].dropna(
            subset=[f"{var}_Lag1", f"{var}_Lag2", "GDP_Growth", "Fed_Funds_Rate", f"{var}_target"])
        if len(train_rows) < 30:
            models[var] = None
            continue
        X = train_rows[[f"{var}_Lag1", f"{var}_Lag2", "GDP_Growth", "Fed_Funds_Rate"]].values
        y = train_rows[f"{var}_target"].values
        model = lgb.LGBMRegressor(**cfg.LGBM_PARAMS_SMALL)
        model.fit(X, y)
        models[var] = model
    return models


def estimate_macro_ar1(hist_values: np.ndarray):
    v = hist_values[~np.isnan(hist_values)]
    if len(v) < 4:
        return (float(v.mean()) if len(v) else 0.0), None, True
    mu = v.mean()
    x, y = v[:-1] - mu, v[1:] - mu
    denom = np.sum(x ** 2)
    if denom == 0:
        return mu, None, True
    omega = np.sum(x * y) / denom
    return (mu, float(omega), False) if 0.0 <= omega <= 1.0 else (mu, None, True)


def project_macro(last_val, mu, omega, use_flat):
    return last_val if use_flat else mu + omega * (last_val - mu)


def forecast_dynamic_features_path(feat_df: pd.DataFrame, valuation_date_quarter: str,
                                    dyn_models: dict, horizon: int = None) -> pd.DataFrame:
    """Recursively rolls forward the 6 dynamic bank-quality variables AND
    the 5 macro variables for JPM. Real panel data used wherever it
    exists; LightGBM (bank vars) / AR(1) (macro vars) beyond that.
    Identical to V4."""
    horizon = horizon or cfg.FORECAST_HORIZON_QUARTERS

    def qkey(q):
        y, qq = int(q[:4]), int(q[-1])
        return y * 4 + qq

    jpm = feat_df[feat_df["Bank"] == cfg.JPM_NAME].set_index("Quarter")
    last_real_quarter = feat_df["Quarter"].max()
    last_real_key = qkey(last_real_quarter)

    seed_quarter = quarter_add(valuation_date_quarter, -1)
    state = {var: jpm.loc[seed_quarter, var] for var in cfg.DYNAMIC_FEATURES}
    state_lag1 = {var: jpm.loc[quarter_add(seed_quarter, -1), var] if quarter_add(seed_quarter, -1) in jpm.index else state[var]
                  for var in cfg.DYNAMIC_FEATURES}
    for mv in cfg.MACRO_FEATURES:
        state[mv] = jpm.loc[seed_quarter, mv]

    jpm_hist = feat_df[(feat_df["Bank"] == cfg.JPM_NAME) & (feat_df["Quarter"] < valuation_date_quarter)].sort_values("Quarter")
    macro_params = {mv: estimate_macro_ar1(jpm_hist[mv].values) for mv in cfg.MACRO_FEATURES}

    rows = []
    for h in range(1, horizon + 1):
        cq = quarter_add(valuation_date_quarter, h - 1)
        this_key = qkey(cq)
        row = {"quarter_ahead": h, "calendar_quarter": cq}

        if this_key <= last_real_key and cq in jpm.index:
            for var in cfg.DYNAMIC_FEATURES:
                new_val = float(jpm.loc[cq, var])
                row[f"{var}_source"] = "actual"
                state_lag1[var] = state[var]
                state[var] = new_val
            for mv in cfg.MACRO_FEATURES:
                state[mv] = float(jpm.loc[cq, mv])
        else:
            for var in cfg.DYNAMIC_FEATURES:
                model = dyn_models.get(var)
                if model is not None:
                    X = np.array([[state[var], state_lag1[var], state["GDP_Growth"], state["Fed_Funds_Rate"]]])
                    new_val = float(model.predict(X)[0])
                else:
                    new_val = state[var]
                row[f"{var}_source"] = "lightgbm_recursive"
                state_lag1[var] = state[var]
                state[var] = new_val
            for mv in cfg.MACRO_FEATURES:
                mu, omega, use_flat = macro_params[mv]
                state[mv] = project_macro(state[mv], mu, omega, use_flat)

        for var in cfg.DYNAMIC_FEATURES:
            row[var] = state[var]
        for mv in cfg.MACRO_FEATURES:
            row[mv] = state[mv]
        rows.append(row)

    return pd.DataFrame(rows)


# ============================================================================
# DIRECT Q1-Q4 MODELS + Q5-Q20 CONTINUATION ENGINE -- identical to V4's
# get_training_rows_as_of / fit_preprocessor / train_direct_horizon_models /
# train_continuation_model / estimate_spread_persistence /
# forecast_roe_spread_path, with ONE parameter added (`omega_cap`) that
# defaults to reproducing V4's own unconstrained behaviour.
# ============================================================================

def get_training_rows_as_of(feat_df: pd.DataFrame, valuation_date_quarter: str, target_col: str) -> pd.DataFrame:
    """No-look-ahead expanding window: target quarter must be strictly
    before the valuation date. Identical to V4."""
    def qkey(q):
        y, qq = int(q[:4]), int(q[-1])
        return y * 4 + qq
    cutoff = qkey(valuation_date_quarter)
    df = feat_df.dropna(subset=cfg.ALL_FEATURES + [target_col]).copy()
    h = int(target_col.split("_")[-1]) if target_col.startswith("ROE_SPREAD_t_plus_") else 1
    df["_qkey"] = df["Quarter"].apply(qkey)
    train_rows = df[df["_qkey"] + h < cutoff].drop(columns=["_qkey"])
    return train_rows.reset_index(drop=True)


def fit_preprocessor(train_df: pd.DataFrame) -> ColumnTransformer:
    pre = ColumnTransformer(transformers=[
        ("num", StandardScaler(), cfg.ALL_FEATURES),
        ("cat", OneHotEncoder(handle_unknown="ignore"), [cfg.CATEGORICAL_FEATURE]),
    ])
    pre.fit(train_df[cfg.ALL_FEATURES + [cfg.CATEGORICAL_FEATURE]])
    return pre


def train_direct_horizon_models(feat_df: pd.DataFrame, valuation_date_quarter: str) -> dict:
    """Trains Model_Q1..Model_Q4, each a SEPARATE LightGBM regressor
    mapping t0's feature vector directly to ROE_SPREAD at t0+h. No
    recursive chaining between them. Identical to V4."""
    models = {}
    for h in range(1, cfg.DIRECT_HORIZON_QUARTERS + 1):
        target_col = f"ROE_SPREAD_t_plus_{h}"
        train_df = get_training_rows_as_of(feat_df, valuation_date_quarter, target_col)
        pre = fit_preprocessor(train_df)
        X = pre.transform(train_df[cfg.ALL_FEATURES + [cfg.CATEGORICAL_FEATURE]])
        y = train_df[target_col].values
        model = lgb.LGBMRegressor(**cfg.LGBM_PARAMS)
        model.fit(X, y)
        models[h] = {"model": model, "preprocessor": pre, "n_train_rows": len(train_df)}
    return models


def train_continuation_model(feat_df: pd.DataFrame, valuation_date_quarter: str) -> dict:
    """The Q5-Q8 continuation engine: ONE pooled one-step-ahead model.
    Identical to V4."""
    features = cfg.CONTINUATION_FEATURES
    train_df = get_training_rows_as_of(feat_df, valuation_date_quarter, "ROE_SPREAD_t_plus_1")
    pre = ColumnTransformer([("num", StandardScaler(), features)])
    pre.fit(train_df[features])
    X = pre.transform(train_df[features])
    y = train_df["ROE_SPREAD"].values
    model = lgb.LGBMRegressor(**cfg.LGBM_PARAMS_SMALL)
    model.fit(X, y)
    return {"model": model, "preprocessor": pre, "features": features, "n_train_rows": len(train_df)}


def estimate_spread_persistence(feat_df: pd.DataFrame, valuation_date_quarter: str) -> float:
    """AR(1)-toward-zero persistence parameter for ROE_SPREAD, estimated
    on JPM's own historical spread series, no look-ahead. This is
    omega_ESTIMATED -- identical calculation to V4. The cap (if any) is
    applied separately downstream, in forecast_roe_spread_path()."""
    def qkey(q):
        y, qq = int(q[:4]), int(q[-1])
        return y * 4 + qq
    cutoff = qkey(valuation_date_quarter)
    jpm_hist = feat_df[(feat_df["Bank"] == cfg.JPM_NAME) & (feat_df["ROE_SPREAD"].notna())
                        & (feat_df["Quarter"].apply(qkey) < cutoff)].sort_values("Quarter")
    spread = jpm_hist["ROE_SPREAD"].values
    if len(spread) < 8:
        return 0.85
    x, y_ = spread[:-1], spread[1:]
    omega_raw = np.sum(x * y_) / np.sum(x ** 2)
    return float(np.clip(omega_raw, 0.0, 1.0))


def forecast_roe_spread_path(feat_df: pd.DataFrame, valuation_date_quarter: str,
                              direct_models: dict, continuation: dict, dyn_models: dict,
                              horizon: int = None, ml_continuation_horizon: int = None,
                              omega_cap=None) -> dict:
    """
    Builds the full 20-quarter ROE_SPREAD path:
      Q1-4:  Model_Q1..Q4, each fed the SAME t0 feature vector (direct).
      Q5-8:  recursive LightGBM continuation, fed by dynamically-forecast
             bank-quality/macro features.
      Q9-20: analytical AR(1)-toward-zero decay.

    `omega_cap` is THE ONLY V4.2 ENHANCEMENT, isolated to one line:
        omega_cap=None  -> omega_final = omega_estimated  (V4's own,
                            UNCONSTRAINED formula -- exactly reproduces V4)
        omega_cap=0.92  -> omega_final = min(omega_estimated, 0.92)
                            (V4.2's constrained formula)
    Every other line in this function is identical to V4's
    forecast_roe_spread_path().
    """
    horizon = horizon or cfg.FORECAST_HORIZON_QUARTERS
    ml_continuation_horizon = ml_continuation_horizon or cfg.ML_CONTINUATION_HORIZON

    seed_quarter = quarter_add(valuation_date_quarter, -1)
    jpm_row = feat_df[(feat_df["Bank"] == cfg.JPM_NAME) & (feat_df["Quarter"] == seed_quarter)].iloc[0]

    spread_path = np.zeros(horizon)
    is_direct = np.zeros(horizon, dtype=bool)
    is_ml_continuation = np.zeros(horizon, dtype=bool)

    # --- Q1-4: direct models, fixed t0 feature vector ---
    X_t0 = pd.DataFrame([{**{f: jpm_row[f] for f in cfg.ALL_FEATURES}, cfg.CATEGORICAL_FEATURE: cfg.JPM_NAME}])[cfg.ALL_FEATURES + [cfg.CATEGORICAL_FEATURE]]
    for h in range(1, cfg.DIRECT_HORIZON_QUARTERS + 1):
        d = direct_models[h]
        X_t = d["preprocessor"].transform(X_t0)
        spread_path[h - 1] = float(d["model"].predict(X_t)[0])
        is_direct[h - 1] = True

    # --- Q5-8: recursive ML continuation ---
    dyn_path = forecast_dynamic_features_path(feat_df, valuation_date_quarter, dyn_models, horizon)
    spread_lag1 = spread_path[cfg.DIRECT_HORIZON_QUARTERS - 1]
    spread_ma4 = np.mean(spread_path[max(0, cfg.DIRECT_HORIZON_QUARTERS - 4):cfg.DIRECT_HORIZON_QUARTERS])

    cmodel, cpre, cfeat = continuation["model"], continuation["preprocessor"], continuation["features"]
    for h in range(cfg.DIRECT_HORIZON_QUARTERS + 1, ml_continuation_horizon + 1):
        dyn_row = dyn_path[dyn_path["quarter_ahead"] == h].iloc[0]
        X_row = {"ROE_SPREAD_Lag1": spread_lag1, "ROE_SPREAD_MA_4Q": spread_ma4}
        for f in cfg.BANK_QUALITY_FEATURES + cfg.MACRO_FEATURES:
            X_row[f] = dyn_row[f]
        X_df = pd.DataFrame([X_row])[cfeat]
        X_t = cpre.transform(X_df)
        pred = float(cmodel.predict(X_t)[0])
        spread_path[h - 1] = pred
        is_ml_continuation[h - 1] = True
        spread_ma4 = np.mean(np.append(spread_path[max(0, h - 4):h - 1], pred))
        spread_lag1 = pred

    # --- Q9-20: analytical AR(1)-toward-zero decay ---
    omega_estimated = estimate_spread_persistence(feat_df, valuation_date_quarter)
    # *** THE ONLY V4.2 MODIFICATION IS THIS ONE LINE: ***
    omega_final = omega_estimated if omega_cap is None else min(omega_estimated, omega_cap)

    spread_prev = spread_path[ml_continuation_horizon - 1]
    for h in range(ml_continuation_horizon + 1, horizon + 1):
        spread_h = omega_final * spread_prev
        spread_path[h - 1] = spread_h
        spread_prev = spread_h

    return {"roe_spread_path_pct": spread_path, "is_direct_horizon": is_direct,
            "is_ml_continuation": is_ml_continuation, "omega_estimated": omega_estimated,
            "omega_cap": omega_cap, "omega_final": omega_final,
            "constraint_bound": (omega_cap is not None) and (omega_estimated > omega_cap),
            "ml_continuation_horizon": ml_continuation_horizon,
            "dynamic_feature_path": dyn_path, "seed_quarter": seed_quarter, "jpm_seed_row": jpm_row}


# ============================================================================
# ROE RECOVERY, RIV MECHANICS, TERMINAL VALUE, INTRINSIC VALUE -- identical
# to V4's recover_roe_path / build_riv_forecast_path / discount_riv_path /
# compute_terminal_value / compute_intrinsic_value / get_actual_market_price.
#
# *** IMPORTANT CORRECTION VS AN EARLIER STANDALONE ATTEMPT ***
# compute_terminal_value() below uses V4's ACTUAL formula:
#     TV_at_horizon = RI_terminal * (1 + g) / (CoE_q - g),   g = 0.0 default
# This does NOT take omega as an argument and never did in V4. omega's
# only effect on terminal value is INDIRECT, via its effect on the Q9-Q20
# spread path (and hence RI_20, the input to this formula). An earlier
# implementation attempt used a different, omega-dependent formula here;
# that was verified WRONG by direct inspection of V4's source code and is
# corrected here.
# ============================================================================

def recover_roe_path(spread_path_pct: np.ndarray, cost_of_equity_pct: float) -> np.ndarray:
    """ROE = ROE_SPREAD + CoE. CoE held constant across the 20-quarter
    horizon. Identical to V4."""
    return spread_path_pct + cost_of_equity_pct


def estimate_dividend_payout_ratio(financials_df: pd.DataFrame, valuation_date: pd.Timestamp, lookback: int = 8) -> float:
    hist = financials_df[financials_df["period_end"] <= valuation_date].tail(lookback).copy()
    hist["eps"] = hist["net_income_avail_common"] / hist["shares_outstanding_mm"]
    hist = hist[hist["eps"] > 0]
    if hist.empty:
        return 0.0
    return float(np.clip((hist["dps"] / hist["eps"]).mean(), 0.0, 1.0))


def build_riv_forecast_path(financials_df: pd.DataFrame, valuation_date: pd.Timestamp,
                             roe_path_pct: np.ndarray, cost_of_equity: float,
                             horizon: int = None) -> pd.DataFrame:
    """Clean-surplus book value roll-forward + residual income:
    RI = BeginningBV * (ROE - CoE) / 4. Identical to V4."""
    horizon = horizon or cfg.FORECAST_HORIZON_QUARTERS
    hist = financials_df[financials_df["period_end"] <= valuation_date].dropna(subset=["roe_quarterly_annualized_pct"]).sort_values("period_end")
    last_bv = hist.iloc[-1]["common_equity"]
    last_shares = hist.iloc[-1]["shares_outstanding_mm"]
    payout = estimate_dividend_payout_ratio(financials_df, valuation_date)

    rows = []
    bv_prev = last_bv
    for h in range(1, horizon + 1):
        roe_h = roe_path_pct[h - 1] / 100.0
        ni_h = (roe_h / 4) * bv_prev
        div_h = payout * ni_h
        bv_h = bv_prev + ni_h - div_h
        ri_h = ni_h - cost_of_equity / 4 * bv_prev
        rows.append({"quarter_ahead": h, "roe_pct": roe_path_pct[h - 1], "book_value_begin": bv_prev,
                     "net_income": ni_h, "dividend": div_h, "book_value_end": bv_h, "residual_income": ri_h})
        bv_prev = bv_h
    out = pd.DataFrame(rows)
    out["payout_ratio"] = payout
    out["shares_outstanding_mm"] = last_shares
    return out


def discount_riv_path(riv_path: pd.DataFrame, cost_of_equity: float) -> pd.DataFrame:
    df = riv_path.copy()
    coe_q = cost_of_equity / 4
    df["discount_factor"] = 1 / (1 + coe_q) ** df["quarter_ahead"]
    df["pv_residual_income"] = df["residual_income"] * df["discount_factor"]
    return df


def compute_terminal_value(riv_path_disc: pd.DataFrame, cost_of_equity: float, g: float = 0.0) -> dict:
    """
    TV = RI_terminal x (1 + g) / (CoE_q - g), at g=0.0 by default.
    THIS IS V4'S ACTUAL FORMULA -- it does NOT depend on omega. omega's
    effect on terminal value flows entirely through RI_terminal (which is
    itself a function of the Q9-Q20 spread path, which IS controlled by
    omega_final upstream in forecast_roe_spread_path).
    """
    coe_q = cost_of_equity / 4
    last = riv_path_disc.iloc[-1]
    ri_terminal = last["residual_income"]
    tv_at_horizon = ri_terminal * (1 + g) / (coe_q - g)
    pv_tv = tv_at_horizon * last["discount_factor"]
    return {"terminal_growth_g": g, "ri_terminal": ri_terminal, "terminal_value_at_horizon": tv_at_horizon,
            "discount_factor_terminal": last["discount_factor"], "pv_terminal_value": pv_tv}


def compute_intrinsic_value(financials_df: pd.DataFrame, valuation_date: pd.Timestamp,
                             riv_path_disc: pd.DataFrame, terminal: dict) -> dict:
    """IntrinsicValue = BookValue(t0) + PV(ExplicitRI) + PV(TerminalValue).
    Identical to V4."""
    hist = financials_df[financials_df["period_end"] <= valuation_date].sort_values("period_end")
    bv0 = hist.iloc[-1]["common_equity"]
    shares = hist.iloc[-1]["shares_outstanding_mm"]
    pv_ri_sum = riv_path_disc["pv_residual_income"].sum()
    pv_tv = terminal["pv_terminal_value"]
    intrinsic_eq = bv0 + pv_ri_sum + pv_tv
    return {"valuation_date": valuation_date, "book_value_t0": bv0, "shares_outstanding_mm": shares,
            "pv_explicit_ri": pv_ri_sum, "pv_terminal_value": pv_tv, "intrinsic_equity_value": intrinsic_eq,
            "intrinsic_value_per_share": intrinsic_eq / shares}


def get_actual_market_price(jpm_prices: pd.DataFrame, valuation_date: pd.Timestamp) -> float:
    avail = jpm_prices[jpm_prices["week_end_date"] <= valuation_date]
    return float(avail.iloc[-1]["jpm_price"])


# ============================================================================
# TRADITIONAL RIV BASELINE -- identical to V4's estimate_traditional_omega /
# traditional_riv_path
# ============================================================================

def estimate_traditional_omega(financials_df: pd.DataFrame, valuation_date: pd.Timestamp, coe: float, min_obs: int = 8) -> float:
    hist = financials_df[financials_df["period_end"] <= valuation_date].dropna(subset=["roe_quarterly_annualized_pct"]).reset_index(drop=True)
    roe = hist["roe_quarterly_annualized_pct"].values / 100.0
    x, y = roe[:-1] - coe, roe[1:] - coe
    omega_raw = np.sum(x * y) / np.sum(x ** 2)
    return float(np.clip(omega_raw, 0.0, 1.0))


def traditional_riv_path(financials_df: pd.DataFrame, valuation_date: pd.Timestamp, coe: float, horizon: int = None) -> tuple:
    horizon = horizon or cfg.FORECAST_HORIZON_QUARTERS
    omega = estimate_traditional_omega(financials_df, valuation_date, coe)
    hist = financials_df[financials_df["period_end"] <= valuation_date].dropna(subset=["roe_quarterly_annualized_pct"]).sort_values("period_end")
    last_roe = hist.iloc[-1]["roe_quarterly_annualized_pct"]
    roe_path = np.zeros(horizon)
    roe_prev = last_roe / 100.0
    for h in range(horizon):
        roe_h = coe + omega * (roe_prev - coe)
        roe_path[h] = roe_h * 100.0
        roe_prev = roe_h
    return roe_path, omega


# ============================================================================
# MODEL EVALUATION -- identical to V4's evaluate_direct_models_vs_actual /
# compute_accuracy_metrics
# ============================================================================

def evaluate_direct_models_vs_actual(feat_df: pd.DataFrame, valuation_date_quarter: str,
                                      direct_models: dict, continuation: dict, coe_hist: pd.DataFrame) -> pd.DataFrame:
    jpm = feat_df[feat_df["Bank"] == cfg.JPM_NAME].set_index("Quarter")
    seed_quarter = quarter_add(valuation_date_quarter, -1)
    coe_map = coe_hist.set_index("Quarter")["cost_of_equity"]

    rows = []
    X_t0 = pd.DataFrame([{**{f: jpm.loc[seed_quarter, f] for f in cfg.ALL_FEATURES}, cfg.CATEGORICAL_FEATURE: cfg.JPM_NAME}])[cfg.ALL_FEATURES + [cfg.CATEGORICAL_FEATURE]]

    cmodel, cpre, cfeat = continuation["model"], continuation["preprocessor"], continuation["features"]

    for h in range(1, cfg.DIRECT_HORIZON_QUARTERS + 1):
        cq = quarter_add(valuation_date_quarter, h - 1)
        actual_spread = jpm.loc[cq, "ROE_SPREAD"] if cq in jpm.index else np.nan
        actual_roe = jpm.loc[cq, "ROE"] if cq in jpm.index else np.nan

        d = direct_models[h]
        multi_horizon_pred = float(d["model"].predict(d["preprocessor"].transform(X_t0))[0])

        prev_q = quarter_add(cq, -1)
        one_step_pred = np.nan
        if prev_q in jpm.index and not jpm.loc[prev_q, cfeat[2:]].isna().any():
            row_in = {"ROE_SPREAD_Lag1": jpm.loc[prev_q, "ROE_SPREAD"], "ROE_SPREAD_MA_4Q": jpm.loc[prev_q, "ROE_SPREAD_MA_4Q"]}
            for f in cfg.BANK_QUALITY_FEATURES + cfg.MACRO_FEATURES:
                row_in[f] = jpm.loc[prev_q, f]
            X_os = pd.DataFrame([row_in])[cfeat]
            one_step_pred = float(cmodel.predict(cpre.transform(X_os))[0])

        coe_pct = coe_map.get(cq, np.nan) * 100
        rows.append({
            "valuation_label": valuation_date_quarter, "quarter_ahead": h, "calendar_quarter": cq,
            "actual_roe_spread": actual_spread, "multi_horizon_pred_spread": multi_horizon_pred,
            "one_step_pred_spread": one_step_pred,
            "actual_roe": actual_roe, "multi_horizon_pred_roe": multi_horizon_pred + coe_pct,
            "one_step_pred_roe": one_step_pred + coe_pct if not np.isnan(one_step_pred) else np.nan,
        })
    return pd.DataFrame(rows)


def compute_accuracy_metrics(df: pd.DataFrame, actual_col: str, pred_col: str) -> dict:
    sub = df.dropna(subset=[actual_col, pred_col])
    if len(sub) == 0:
        return {"n": 0, "mae": np.nan, "rmse": np.nan, "mape": np.nan}
    e = sub[actual_col] - sub[pred_col]
    return {"n": len(sub), "mae": e.abs().mean(), "rmse": np.sqrt((e ** 2).mean()),
            "mape": (e.abs() / sub[actual_col].abs()).mean() * 100}


# ============================================================================
# VALIDATION SUITE (13 mandatory checks). Each function returns
# (passed: bool, detail: str). The orchestration script (enhanced_riv_v4_2.py)
# HALTS EXECUTION if any check fails -- per the brief's mandatory
# requirement that the program "must stop with an error if any validation
# fails".
# ============================================================================

def validate_cost_of_equity_identical(coe_hist: pd.DataFrame) -> tuple:
    """Validation 1: re-derives CoE independently for every valuation date
    and checks it matches the value stored in coe_hist exactly (proves the
    single CAPM code path is being used consistently, not two diverging
    implementations)."""
    tbill = load_tbill_quarterly()
    jpm_prices = load_jpm_weekly_prices()
    sp500 = load_sp500_daily()
    rp = build_return_panel(jpm_prices, sp500)
    max_diff = 0.0
    for _, row in coe_hist.dropna(subset=["cost_of_equity"]).iterrows():
        as_of = row["quarter_end"]
        recomputed = estimate_cost_of_equity(tbill, rp, as_of)
        max_diff = max(max_diff, abs(recomputed["cost_of_equity"] - row["cost_of_equity"]))
    return max_diff < 1e-9, f"Max independent-recomputation CoE difference: {max_diff:.2e}"


def validate_feature_list(features_used: list) -> tuple:
    """Validation 2: feature list must be exactly 20, matching the brief's
    exact specification (already enforced by asserts in config.py at
    import time, re-checked here for the report)."""
    expected = set(cfg.ALL_FEATURES)
    actual = set(features_used)
    ok = (len(cfg.ALL_FEATURES) == 20) and (expected == actual)
    return ok, f"Feature count: {len(cfg.ALL_FEATURES)} (expected 20). Sets match: {expected == actual}"


def validate_training_row_counts(all_dates_data: list) -> tuple:
    """Validation 3: training row counts must strictly increase across
    the 9 rolling valuation dates (proves the expanding window is genuinely
    expanding, no look-ahead)."""
    counts = [d["n_train_rows_q1"] for d in all_dates_data]
    increasing = all(counts[i] < counts[i + 1] for i in range(len(counts) - 1))
    return increasing, f"Model_Q1 training rows by date: {counts}"


def validate_feature_order(direct_models: dict) -> tuple:
    """Validation 4: the ColumnTransformer's output feature order must be
    identical across all 4 horizon models (same preprocessor recipe)."""
    orders = [list(direct_models[h]["preprocessor"].get_feature_names_out()) for h in direct_models]
    ok = all(o == orders[0] for o in orders)
    return ok, f"All {len(orders)} direct models share an identical post-transform feature order: {ok}"


def validate_q1_q4_unbound_match(all_dates_data: list) -> tuple:
    """Validation 5: Q1-Q4 forecasts must be identical between the capped
    and uncapped path (verified numerically, not assumed -- the cap is
    only ever applied in the Q9-Q20 phase)."""
    max_diff = 0.0
    for d in all_dates_data:
        diff = np.abs(d["spread_path_capped"][:cfg.DIRECT_HORIZON_QUARTERS] -
                      d["spread_path_uncapped"][:cfg.DIRECT_HORIZON_QUARTERS]).max()
        max_diff = max(max_diff, diff)
    return max_diff < 1e-9, f"Max |Q1-Q4 spread difference| (capped vs uncapped), all dates: {max_diff:.2e}"


def validate_q5_q8_unbound_match(all_dates_data: list) -> tuple:
    """Validation 6: Q5-Q8 forecasts must also be identical between the
    capped and uncapped path (the cap only affects Q9-Q20)."""
    max_diff = 0.0
    for d in all_dates_data:
        s, e = cfg.DIRECT_HORIZON_QUARTERS, cfg.ML_CONTINUATION_HORIZON
        diff = np.abs(d["spread_path_capped"][s:e] - d["spread_path_uncapped"][s:e]).max()
        max_diff = max(max_diff, diff)
    return max_diff < 1e-9, f"Max |Q5-Q8 spread difference| (capped vs uncapped), all dates: {max_diff:.2e}"


def validate_book_value_reconciliation(riv_path: pd.DataFrame) -> tuple:
    """Validation 7: EndingBV_t must equal BeginningBV_(t+1) (continuity)."""
    end_prev = riv_path["book_value_end"].shift(1)
    begin_curr = riv_path["book_value_begin"]
    max_dev = (end_prev - begin_curr).abs().dropna().max()
    return max_dev < 1e-6, f"Max book value continuity deviation: {max_dev:.2e}"


def validate_clean_surplus(riv_path: pd.DataFrame) -> tuple:
    """Validation 8: BeginningBV + NI - Dividends = EndingBV exactly."""
    implied_end = riv_path["book_value_begin"] + riv_path["net_income"] - riv_path["dividend"]
    max_dev = (implied_end - riv_path["book_value_end"]).abs().max()
    return max_dev < 1e-6, f"Max clean surplus deviation: {max_dev:.2e}"


def validate_discount_factors(riv_disc: pd.DataFrame) -> tuple:
    """Validation 9: discount factors strictly decreasing, in (0,1)."""
    dfs = riv_disc["discount_factor"].values
    ok = bool(np.all(np.diff(dfs) < 0) and dfs[0] < 1.0 and dfs[0] > 0)
    return ok, f"Discount factors strictly decreasing and in (0,1): {ok}"


def validate_terminal_denominator(coe: float, g: float = 0.0) -> tuple:
    """Validation 10: terminal value denominator (CoE_q - g) must be
    strictly positive for the perpetuity formula to be well-defined."""
    coe_q = coe / 4
    ok = (coe_q - g) > 0
    return ok, f"CoE_q - g = {coe_q - g:.6f} > 0: {ok}"


def validate_intrinsic_value_reconciliation(iv: dict) -> tuple:
    """Validation 11: BV(t0) + PV(ExplicitRI) + PV(TerminalValue) must
    equal the reported intrinsic equity value exactly."""
    implied = iv["book_value_t0"] + iv["pv_explicit_ri"] + iv["pv_terminal_value"]
    dev = abs(implied - iv["intrinsic_equity_value"])
    return dev < 1e-3, f"Reconciliation deviation: {dev:.2e}"


def validate_v4_2_equals_v4_when_unbound(all_dates_data: list) -> tuple:
    """Validation 12 (the key V4.2 requirement): for every date where
    omega_estimated <= OMEGA_CAP, the capped (V4.2) and uncapped
    (V4-equivalent) intrinsic values -- built from the SAME trained
    models and the SAME upstream path -- must be identical to
    floating-point precision."""
    max_diff = 0.0
    for d in all_dates_data:
        if not d["constraint_bound"]:
            max_diff = max(max_diff, abs(d["iv_capped"] - d["iv_uncapped"]))
    return max_diff < 1e-6, f"Max |IV(V4.2) - IV(V4)| across all dates with omega_estimated<=0.92: {max_diff:.2e}"


def validate_only_bound_dates_differ(all_dates_data: list) -> tuple:
    """Validation 13: dates where omega_estimated > OMEGA_CAP must show a
    genuine, non-zero difference (proves the cap actually does something
    where it should, not just that it does nothing where it shouldn't)."""
    bound_dates = [d for d in all_dates_data if d["constraint_bound"]]
    unbound_dates = [d for d in all_dates_data if not d["constraint_bound"]]
    bound_all_differ = all(abs(d["iv_capped"] - d["iv_uncapped"]) > 1e-9 for d in bound_dates) if bound_dates else True
    unbound_all_match = all(abs(d["iv_capped"] - d["iv_uncapped"]) < 1e-9 for d in unbound_dates)
    ok = bound_all_differ and unbound_all_match
    return ok, (f"{len(bound_dates)} date(s) with omega>0.92 all differ: {bound_all_differ}. "
                f"{len(unbound_dates)} date(s) with omega<=0.92 all match: {unbound_all_match}.")
